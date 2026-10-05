#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
接続検討回答の検証一覧（xlsx）を、変電所スコアの実績DBに取り込むための JSON に変換する。

入力: 「接続検討回答 検証一覧」xlsx の「回答一覧」シート（回答書1件=1行）
出力: data/connection_results.json
      substation_scoring.py が読み込み、手書きの実績DB（RESULTS_DB）の後ろに追加する。
      → score_substations.py（xlsx版）と build_substation_web_data.py（Web版）の両方に反映される。

取り込む行:
  - 「集計対象」が ○ の行（案件ごとの最新の回答で、負担金の記載があるもの）
  - ただし手書きの RESULTS_DB に同じ案件が既にある行（SKIP_NOS）は除く

変電所名は全国変電所リストの表記に合わせる（Web版はエリア＋変電所名で引き当てるため）。
回答書の「変電所」欄の先頭がリストの変電所名と一致する場合だけ採用し、
それ以外（基幹変電所しか書かれていない・「本文に記載なし」等）は「－(…)」と原文を残す。

使い方:
  python scripts/import_connection_results.py \
    --in "data/接続検討回答_検証一覧.xlsx" \
    --list "data/全国変電所リスト_66kV以上_座標追加.xlsx" \
    --out data/connection_results.json
"""
import argparse
import datetime
import json
import re
import unicodedata

import openpyxl

from score_substations import read_substations

# 手書きの RESULTS_DB（substation_scoring.py）に既に載っている案件の行（回答一覧のNo）。
# 11=伊勢市特高 / 40-42=熊本県南関町(西日本PE) / 75=喜多方市松山(No071)
SKIP_NOS = {11, 40, 41, 42, 75}

COLS = ["No", "案件番号", "案件名", "回答種別", "電力会社", "回答日", "放電kW", "電圧", "変電所",
        "負担金(税込・万円)", "保証金(円)", "工期(月)", "工期の原文", "主な対策工事", "上位系統増強",
        "読取信頼度", "集計対象", "律速・可否メモ"]


def nfkc(s):
    # 「駒ケ嶺」と「駒ヶ嶺」のような表記ゆれを揃える
    return unicodedata.normalize("NFKC", str(s or "")).strip().replace("ケ", "ヶ")


def man(v):
    """万円の表記（1万円未満は四捨五入）。"""
    return f"{v:,.0f}万円" if v < 10000 else f"{v / 10000:.2f}億円"


def cut(s, n):
    s = re.sub(r"\s+", " ", str(s or "")).strip()
    return s if len(s) <= n else s[: n - 1] + "…"


def match_substation(text, names):
    """回答書の変電所欄の先頭がリストの変電所名に一致すればその名前を返す（最長一致）。

    names は {照合用の名前: リスト表記} 。「真加部変電所の近隣」のように
    接続先そのものではない書き方は採用しない。
    """
    t = nfkc(text)
    best = ""
    for n in names:
        if t.startswith(n) and not t[len(n):].startswith("の") and len(n) > len(best):
            best = n
    return names[best] if best else ""


def read_rows(path):
    ws = openpyxl.load_workbook(path, data_only=True)["回答一覧"]
    header = [c.value for c in ws[1]]
    idx = {h: header.index(h) for h in COLS}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[idx["No"]] is None:
            continue
        yield {h: r[i] for h, i in idx.items()}


def to_record(row, area_names):
    area = row["電力会社"]
    raw = row["変電所"] or ""
    name = match_substation(raw, area_names.get(area, {}))
    if not name:
        name = f"－({cut(raw, 30)})" if raw else "－"
    site = re.sub(r"^\d+_", "", str(row["案件名"] or ""))
    kw = row["放電kW"]
    out = f"{kw:,.0f}kW" if isinstance(kw, (int, float)) else (row["電圧"] or "")
    cost_v = row["負担金(税込・万円)"]
    cost = f"{man(cost_v)}(税込)" if isinstance(cost_v, (int, float)) else "記載なし"
    dep = row["保証金(円)"]
    if isinstance(dep, (int, float)) and dep:
        cost += f"+保証金{dep / 10000:,.1f}万"
    months = row["工期(月)"]
    term = f"入金後{months:.0f}ヶ月" if isinstance(months, (int, float)) else cut(row["工期の原文"], 40) or "記載なし"
    works = cut(row["主な対策工事"], 80)
    upper = cut(row["上位系統増強"], 60)
    lesson = "／".join(x for x in [row["律速・可否メモ"], f"上位系: {upper}" if upper else ""] if x)
    if row["読取信頼度"] != "high":
        lesson += f"（AI読取・信頼度{row['読取信頼度']}・要原本確認）"
    date = row["回答日"] or "日付不明"
    short_cost = man(cost_v) if isinstance(cost_v, (int, float)) else "負担金記載なし"
    short_term = f"{months:.0f}ヶ月" if isinstance(months, (int, float)) else "工期記載なし"
    summary = f"{row['回答種別']}: {short_cost}・{short_term}({date})"
    if row["律速・可否メモ"]:
        summary += f" {row['律速・可否メモ']}"
    return [area, name, site, out, cost, term, works, lesson, summary]


def main():
    ap = argparse.ArgumentParser(description="接続検討回答の検証一覧を実績DB用JSONに変換")
    ap.add_argument("--in", dest="src", required=True, help="接続検討回答 検証一覧のxlsx")
    ap.add_argument("--list", dest="lst", required=True, help="全国変電所リスト（66kV以上）のxlsx")
    ap.add_argument("--list-sheet", default="全国変電所リスト")
    ap.add_argument("--out", required=True, help="出力JSON（例: data/connection_results.json）")
    args = ap.parse_args()

    area_names = {}
    for row in read_substations(args.lst, args.list_sheet):
        area, name = row[1], (row[2] or "").strip()
        if name:
            area_names.setdefault(area, {})[nfkc(name)] = name

    records, matched = [], 0
    for row in read_rows(args.src):
        if row["集計対象"] != "○" or int(row["No"]) in SKIP_NOS:
            continue
        rec = to_record(row, area_names)
        matched += not rec[1].startswith("－")
        records.append(rec)

    with open(args.out, "w", encoding="utf-8") as f:
        json.dump({"generated": datetime.date.today().isoformat(),
                   "source": "接続検討回答 検証一覧（回答一覧シート・集計対象○の行）",
                   "records": records}, f, ensure_ascii=False, indent=1)
    print(f"{len(records)}件を出力（変電所名がリストと一致: {matched}件）→ {args.out}")


if __name__ == "__main__":
    main()
