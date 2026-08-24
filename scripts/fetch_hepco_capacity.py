#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
北海道電力ネットワークの系統空容量CSVから、本体xlsxの北海道行に容量情報を記入する。

北海道は2026-07の全国反映時に設備容量・運用容量・N-1・出力制御が公式リスト由来で
一部埋まったが、予想潮流・空容量（当該/上位系）は全385行が空欄のまま残っていた。
HEPCOは基幹＋ローカル7エリアの変圧器一覧をCSV(ZIP)で公表しており、
本体xlsxと同じ列構成（予想潮流・空容量・N-1電制・出力制御）を持つため、
名称照合でそのまま反映できる。

  https://www.hepco.co.jp/network/con_service/public_document/bid_info.html
  zip/sys_capa_kikan.zip, zip/sys_capa_local01.zip 〜 local07.zip

【照合ルール】取り違えはスコアの誤りになるため、一意に決まるときだけ採用する
  1. 正規化した変電所名で照合（両者とも北海道のみなので名称衝突は同一変電所の
     バンク違いだけ。fetch_substation_coords.normalize_name を流用）
  2. 候補が複数（＝複数バンク）なら、本体xlsxの二次電圧と一致するバンクに絞る
  3. まだ複数なら最大電圧（一次）で絞る
  4. それでも一意にならなければ見送り、レポートに残す

【書き込みルール】空欄のみ埋める。既存値との差異は上書きせずレポートに記録する
  - 予想潮流・空容量(当該/上位系): CSVの値（"―"は空欄のまま）
  - 設備容量・運用容量・二次電圧・N-1電制: 空欄のときだけ
  - 出力制御: 空欄かつCSVが「有り」「なし」のときだけ（"―"は不明なので書かない）

使い方:
  # ① 取得（要ネットワーク）。ZIPを hepco_cache/ に展開する
  python scripts/fetch_hepco_capacity.py fetch --cache-dir hepco_cache

  # ② 反映（ネットワーク不要）
  python scripts/fetch_hepco_capacity.py apply \
    --in "data/全国変電所リスト_66kV以上_座標追加.xlsx" --cache-dir hepco_cache \
    --out "data/全国変電所リスト_66kV以上_座標追加.xlsx" --report hepco_capacity_report.csv
"""
import argparse
import csv
import glob
import io
import os
import re
import sys
import urllib.request
import zipfile

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from fetch_substation_coords import normalize_name, read_rows

BASE = "https://www.hepco.co.jp/network/con_service/public_document/zip/"
ZIPS = ["sys_capa_kikan.zip"] + [f"sys_capa_local{i:02d}.zip" for i in range(1, 8)]
FILL = PatternFill("solid", fgColor="DDEBF7")   # 青=公表資料からの追記

# CSV列 → 本体xlsxの列（read_rows のヘッダートークン）
CSV_COLS = {
    "vsec": "電圧(二次)(kV)",
    "cap": "設備容量(100%×台数)(MW)",
    "opcap": "運用容量値(MW)",
    "flow": "予想潮流(MW)",
    "avail": "空容量(当該設備)(MW)",
    "avail_up": "空容量(上位系等考慮)(MW)",
    "n1": "N-1電制適用可否",
    "curtail": "平常時出力制御の可能性",
}
XLSX_TOKENS = {
    "vsec": "二次電圧", "cap": "設備容量", "opcap": "運用容量", "flow": "予想潮流",
    "avail": "当該設備", "avail_up": "上位系考慮", "n1": "N-1電制", "curtail": "出力制御",
}


def num(s):
    s = str(s or "").strip().replace(",", "")
    if s in ("", "―", "-", "－", "—"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def clean(s):
    """全角空白を半角に潰す（N-1「不可　♯１」→「不可 ♯１」）。"""
    return re.sub(r"\s+", " ", str(s or "").replace("　", " ")).strip()


def cmd_fetch(args):
    os.makedirs(args.cache_dir, exist_ok=True)
    for name in ZIPS:
        req = urllib.request.Request(BASE + name, headers={"User-Agent": "Mozilla/5.0"})
        data = urllib.request.urlopen(req, timeout=120).read()
        n = 0
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            for info in z.infolist():
                # ZIP内のファイル名はcp932。変圧器(Tr)のみ使う
                fn = info.filename.encode("cp437").decode("cp932", "replace")
                if "_Tr_" not in fn:
                    continue
                out = os.path.join(args.cache_dir, os.path.basename(fn))
                with open(out, "wb") as f:
                    f.write(z.read(info))
                n += 1
        print(f"{name}: 変圧器CSV {n}件")
    print(f"完了: {args.cache_dir}")


def load_csv_index(cache_dir):
    """{正規化名: [バンク行dict]} を返す。"""
    paths = sorted(glob.glob(os.path.join(cache_dir, "*_Tr_*.csv")))
    if not paths:
        raise SystemExit(f"変圧器CSVがありません: {cache_dir}（先に fetch を実行）")
    idx = {}
    for p in paths:
        rows = list(csv.reader(open(p, encoding="cp932")))
        hdr_i = next(i for i, r in enumerate(rows) if r and "変電所名" in [c.strip() for c in r])
        col = {}
        for i, h in enumerate(rows[hdr_i]):
            h = h.strip()
            for key, token in CSV_COLS.items():
                # 実ファイルの列名は改行等の揺れがあるため前方一致で拾う
                if h.startswith(token.split("(")[0]) and token.split("(")[0] and key not in col \
                        and h.replace(" ", "").startswith(token.replace(" ", "")[:6]):
                    col[key] = i
        col["name"] = next(i for i, h in enumerate(rows[hdr_i]) if h.strip() == "変電所名")
        col["v1"] = next(i for i, h in enumerate(rows[hdr_i]) if h.strip().startswith("電圧(一次)"))
        col["v2"] = next(i for i, h in enumerate(rows[hdr_i]) if h.strip().startswith("電圧(二次)"))
        for r in rows[hdr_i + 1:]:
            if len(r) <= col["name"] or not r[col["name"]].strip():
                continue
            bank = {
                "name": r[col["name"]].strip(),
                "src": os.path.basename(p),
                "v1": num(r[col["v1"]]),
                "v2": num(r[col["v2"]]),
            }
            for key in ("cap", "opcap", "flow", "avail", "avail_up"):
                bank[key] = num(r[col[key]]) if key in col else None
            bank["n1"] = clean(r[col["n1"]]) if "n1" in col else ""
            bank["curtail"] = clean(r[col["curtail"]]) if "curtail" in col else ""
            key = normalize_name(bank["name"])
            if key:
                idx.setdefault(key, []).append(bank)
    return idx


def pick(cands, vmax, vsec):
    """バンク候補から一意に決める。決められなければ None。"""
    if len(cands) == 1:
        return cands[0]
    if isinstance(vsec, (int, float)):
        c = [b for b in cands if b["v2"] == vsec]
        if len(c) == 1:
            return c[0]
        if c:
            cands = c
    if isinstance(vmax, (int, float)):
        c = [b for b in cands if b["v1"] == vmax]
        if len(c) == 1:
            return c[0]
    return None


def cmd_apply(args):
    idx = load_csv_index(args.cache_dir)
    rows, cols, _ = read_rows(args.src, args.list_sheet)

    wb = load_workbook(args.src)
    ws = wb[args.list_sheet]
    # 書き込み先の列番号（read_rows は座標系の列しか返さないため自前で引く）
    hdr_row = next(r[0].row for r in ws.iter_rows(min_row=1, max_row=10) if r[0].value == "No.")
    headers = [str(c.value or "").replace("\n", "") for c in ws[hdr_row]]
    xcol = {}
    for key, token in XLSX_TOKENS.items():
        xcol[key] = next(i + 1 for i, h in enumerate(headers) if token in h)

    filled_rows, skipped, report = 0, {"一致なし": 0, "候補複数で特定不可": 0}, []
    for r in rows:
        if r["area"] != "北海道":
            continue
        cands = idx.get(normalize_name(r["name"]), [])
        if not cands:
            skipped["一致なし"] += 1
            report.append([r["no"], r["name"], "一致なし", "", "", ""])
            continue
        vsec = ws.cell(r["row"], xcol["vsec"]).value   # read_rowsは二次電圧を返さない
        bank = pick(cands, r["vmax"], vsec)
        if bank is None:
            skipped["候補複数で特定不可"] += 1
            report.append([r["no"], r["name"], f"候補{len(cands)}件で特定不可",
                           "", "", "; ".join(f'{b["v1"]:g}/{b["v2"]:g}' for b in cands)])
            continue

        wrote, diffs = [], []
        for key in ("vsec", "cap", "opcap", "flow", "avail", "avail_up"):
            new = bank["v2"] if key == "vsec" else bank.get(key)
            if new is None:
                continue
            cell = ws.cell(r["row"], xcol[key])
            if cell.value in (None, ""):
                cell.value = new
                cell.fill = FILL
                wrote.append(key)
            elif isinstance(cell.value, (int, float)) and float(cell.value) != new:
                diffs.append(f"{key}: 既存{cell.value:g}→CSV{new:g}")
        for key in ("n1", "curtail"):
            new = bank.get(key, "")
            if key == "curtail" and new not in ("有り", "なし"):
                continue
            if not new or num(new) is None and new in ("―", "-"):
                continue
            cell = ws.cell(r["row"], xcol[key])
            if cell.value in (None, ""):
                cell.value = new
                cell.fill = FILL
                wrote.append(key)
        if wrote:
            filled_rows += 1
        report.append([r["no"], r["name"], f"一致({bank['v1']:g}/{bank['v2']:g}kV)",
                       ",".join(wrote), "; ".join(diffs), bank["src"]])

    wb.save(args.out)
    print(f"北海道: 記入 {filled_rows}行 / 一致なし {skipped['一致なし']}行 / "
          f"特定不可 {skipped['候補複数で特定不可']}行")
    print(f"出力: {args.out}")
    if args.report:
        with open(args.report, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f)
            w.writerow(["No.", "変電所名", "照合結果", "記入した列", "既存値との差異(未上書き)", "CSV/バンク"])
            w.writerows(report)
        print(f"レポート: {args.report}")
    print("※ 一致しなかった行は空欄のままです（HEPCO非公表の変電所）。")
    print("※ 反映後は score_substations.py と build_substation_web_data.py の再実行が必要です。")


def main():
    ap = argparse.ArgumentParser(description="HEPCOの空容量CSVを北海道行に反映する")
    sub = ap.add_subparsers(dest="cmd", required=True)
    f = sub.add_parser("fetch", help="CSV(ZIP)をダウンロードして展開（要ネットワーク）")
    f.add_argument("--cache-dir", required=True)
    f.set_defaults(func=cmd_fetch)
    a = sub.add_parser("apply", help="CSVを本体xlsxへ反映（ネットワーク不要）")
    a.add_argument("--in", dest="src", required=True)
    a.add_argument("--cache-dir", required=True)
    a.add_argument("--out", dest="out", required=True)
    a.add_argument("--report")
    a.add_argument("--list-sheet", default="全国変電所リスト")
    a.set_defaults(func=cmd_apply)
    args = ap.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
