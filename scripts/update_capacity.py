#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
スコアの元データ（空容量・運用容量・予想潮流・N-1電制・出力制御）を、各送配電事業者の
最新の公表CSVで更新する。月1回の定期更新（.github/workflows/capacity-refresh.yml）で使う。

対象は、全国共通形式（OCCTOの系統空容量公表フォーマット）の変電所CSVを公開している6社と、
同じ項目をPDF（混雑状況マップ・予想潮流等一覧表）で公表している北海道。
  東北・北陸・関西・中国・四国・九州・北海道（PDFは hokkaido_capacity_pdf.py で読む。要 pdfplumber）
東京（Akamaiが全パスに403）・中部（Azure Front Doorが403）・沖縄（自動取得不可）
は対象外で、レポートに「手作業で確認」と出す。

  # ① 取得：各社のページから最新の変圧器CSV（ZIP含む）を探してダウンロード（要ネットワーク）
  python scripts/update_capacity.py fetch --cache-dir capacity_cache          # 全エリア
  python scripts/update_capacity.py fetch --cache-dir capacity_cache --areas 北海道   # 一部だけ

  # ② 反映：本体xlsxの該当行の容量欄を書き換える（ネットワーク不要）
  python scripts/update_capacity.py apply \
    --in "data/全国変電所リスト_66kV以上_座標追加.xlsx" --cache-dir capacity_cache \
    --out "data/全国変電所リスト_66kV以上_座標追加.xlsx" --report capacity_report.csv

【照合】エリア内で、変電所名（「変電所」などの種別を外した名前）が一致する行を探す。
  同じ変電所に複数のバンク（二次電圧違い）がある場合は、本体の二次電圧・最大電圧で1つに絞る。
  絞れなければ書き換えずに「特定不可」としてレポートに残す（推測で選ばない）。
  CSVにあって本体に無い変電所は追加しない（座標などが無いため）。件数だけレポートする。
"""
import argparse
import csv
import datetime
import io
import json
import os
import re
import sys
import time
import unicodedata
import urllib.parse
import urllib.request
import zipfile

from openpyxl import load_workbook

from fetch_substation_coords import loose_name

UA = "Mozilla/5.0 (battery-storage-dd capacity updater; +https://github.com/yonomotozumi-spec/battery-storage-dd)"

# エリア → (公表ページ, 変電所（変圧器）CSV/ZIPのリンクを拾う正規表現)
SOURCES = {
    "東北": ("https://nw.tohoku-epco.co.jp/consignment/system/announcement/index.html",
             r"sys_capa_(?:kikan|local)\d+_tr_\d{6}_\d+\.csv"),
    "北陸": ("https://www.rikuden.co.jp/nw_notification/U_154seiyaku.html",
             r"sys_capa_(?:kikan|local)\d+_tr_\d{6}_\d+\.csv"),
    "関西": ("https://www.kansai-td.co.jp/consignment/disclosure/distribution-equipment/index.html",
             r"154kv_(?:more|less)_trans\.csv"),
    "中国": ("https://www.energia.co.jp/nw/service/retailer/keitou/access/index.html",
             r"zip/csv_[a-z0-9]+\.zip"),
    "四国": ("https://www.yonden.co.jp/nw/line_access/data.html",
             r"sys_capa_(?:kikan|local)\d+_tr_\d{6}_\d+\.csv"),
    # 九州は送電線ZIPと変圧器ZIPが同じ td_*.zip 名で並ぶ。中身のフォルダ名「変圧器CSV」で選ぶ
    "九州": ("https://www.kyuden.co.jp/td/service/wheeling/disclosure.html",
             r"/var/rev0/[0-9/]+/td_[A-Za-z0-9]+\.zip"),
    # 北海道はCSVが無く、系統別PDF（01〜24）と187kV以上PDFの「予想潮流等一覧表」を読む
    "北海道": ("https://www.hepco.co.jp/network/con_service/public_document/bid_info.html",
              r"pdf/map_forecast_tide_list(?:_\d+)?\.pdf"),
}
MANUAL_AREAS = {
    "東京": "www.tepco.co.jp は自動取得を全パスで拒否（403）。県別PDF「予想潮流等一覧表（変電所）」を手作業で確認",
    "中部": "gridmap.powergrid.chuden.co.jp は自動取得を拒否（403）。ブラウザで開き、一覧CSV（KRSIH010〜016）を取得",
    "沖縄": "www.okiden.co.jp は自動取得を拒否（403）。con_res_map*.csv を手作業で確認",
}

DASHES = {"", "-", "－", "ー", "―", "—", "‐", "−", "–"}


def http_get(url, retries=3):
    req = urllib.request.Request(url, headers={"User-Agent": UA})
    for i in range(retries):
        try:
            with urllib.request.urlopen(req, timeout=60) as r:
                return r.read()
        except Exception:  # noqa: BLE001
            if i == retries - 1:
                raise
            time.sleep(3 * (i + 1))


def decode(b):
    for enc in ("utf-8-sig", "cp932"):
        try:
            return b.decode(enc)
        except UnicodeDecodeError:
            continue
    return b.decode("cp932", "replace")


def cache_name(area, url):
    return re.sub(r"[^\w.\-]", "_", f"{area}__{os.path.basename(urllib.parse.urlparse(url).path)}")


def cmd_fetch(args):
    os.makedirs(args.cache_dir, exist_ok=True)
    manifest = {"fetched": datetime.date.today().isoformat(), "areas": {}}
    for area, (page, pat) in SOURCES.items():
        if args.areas and area not in args.areas:
            continue
        try:
            html = decode(http_get(page))
        except Exception as e:  # noqa: BLE001
            print(f"{area}: ページ取得失敗 {e}", file=sys.stderr)
            manifest["areas"][area] = {"error": f"ページ取得失敗: {e}"}
            continue
        # href をページURL基準で解決する（"./data/..." "/nw/..." "zip/..." が混在する）
        links = sorted({urllib.parse.urljoin(page, h) for h in re.findall(r'href="([^"]+)"', html)
                        if re.search(pat, h)})
        files = []
        for url in links:
            try:
                b = http_get(url)
            except Exception as e:  # noqa: BLE001
                print(f"{area}: {url} 取得失敗 {e}", file=sys.stderr)
                continue
            if url.endswith(".zip"):
                try:
                    zf = zipfile.ZipFile(io.BytesIO(b))
                except zipfile.BadZipFile:
                    continue
                names = [n for n in zf.namelist() if n.lower().endswith(".csv")]
                # 変圧器（変電所）のCSVだけ。九州は「変圧器CSV/」フォルダ、中国は *_tr_*.csv
                names = [n for n in names if "変圧器" in n or re.search(r"_tr_", n)]
                if not names:
                    continue
                for n in names:
                    fn = cache_name(area, url) + "__" + re.sub(r"[^\w.\-]", "_", os.path.basename(n))
                    with open(os.path.join(args.cache_dir, fn), "wb") as f:
                        f.write(zf.read(n))
                    files.append({"file": fn, "url": url, "member": n})
            else:
                fn = cache_name(area, url)
                with open(os.path.join(args.cache_dir, fn), "wb") as f:
                    f.write(b)
                files.append({"file": fn, "url": url})
            time.sleep(args.sleep)
        manifest["areas"][area] = {"page": page, "files": files}
        print(f"{area}: {len(files)}ファイル")
    with open(os.path.join(args.cache_dir, "manifest.json"), "w", encoding="utf-8") as f:
        json.dump(manifest, f, ensure_ascii=False, indent=1)


# ---------- CSV の読み取り ----------

def num(s):
    # 0 を空扱いにしないよう `str(s or "")` は使わない（空容量0が書き込まれなくなる）
    s = unicodedata.normalize("NFKC", "" if s is None else str(s)).replace(",", "").replace("'", "").strip()
    if s in DASHES or not re.fullmatch(r"-?\d+(\.\d+)?", s):
        return None
    v = float(s)
    return int(v) if v == int(v) else v


def text(s):
    s = unicodedata.normalize("NFKC", "" if s is None else str(s)).replace("♯", "#").replace("'", "")
    return re.sub(r"\s+", " ", s).strip()


def n1_value(s):
    t = text(s)
    return None if t in DASHES else t


def curtail_value(s):
    # 本体xlsxの既存データは「有り」以外（－・ー）を「なし」として転記している
    t = text(s)
    if not t:
        return None
    return "有り" if t.startswith("有") else "なし"


def find_col(header, *must, avoid=()):
    for i, h in enumerate(header):
        h2 = unicodedata.normalize("NFKC", h).replace(" ", "")
        if all(m in h2 for m in must) and not any(a in h2 for a in avoid):
            return i
    return None


def parse_csv(path):
    """(更新日の文字列, 行のリスト) を返す。行は dict。"""
    t = decode(open(path, "rb").read())
    rows = list(csv.reader(io.StringIO(t)))
    date = text(rows[0][0]) if rows and rows[0] else ""
    hi = next((i for i, r in enumerate(rows[:10]) if any("変電所名" in (c or "") for c in r)), None)
    if hi is None:
        return date, []
    h = [c or "" for c in rows[hi]]
    col = {
        "name": find_col(h, "変電所名"),
        "v1": find_col(h, "一次"),
        "v2": find_col(h, "二次"),
        "cap": find_col(h, "設備容量"),
        "opcap": find_col(h, "運用容量値"),
        "flow": find_col(h, "予想潮流"),
        "cur": find_col(h, "当該設備", "空") if find_col(h, "当該設備", "空") is not None
        else find_col(h, "空き容量"),
        "up": find_col(h, "上位系", "空"),
        "n1": find_col(h, "N-1電制適用可否"),
        "curtail": find_col(h, "平常時出力制御", avoid=("設備", "系統")),
    }
    out = []
    for r in rows[hi + 1:]:
        g = lambda k: (r[col[k]] if col[k] is not None and col[k] < len(r) else "")  # noqa: E731
        name = text(g("name"))
        if not name:
            continue
        out.append({
            "name": name, "v1": num(g("v1")), "v2": num(g("v2")),
            "cap": num(g("cap")), "opcap": num(g("opcap")), "flow": num(g("flow")),
            "cur": num(g("cur")), "up": num(g("up")) if col["up"] is not None else None,
            "n1": n1_value(g("n1")), "curtail": curtail_value(g("curtail")),
            "has_up": col["up"] is not None,
        })
    return date, out


# ---------- 本体xlsxへの反映 ----------

FIELDS = [("cap", "設備容量"), ("opcap", "運用容量"), ("flow", "予想潮流"), ("cur", "当該設備"),
          ("up", "上位系考慮"), ("n1", "N-1電制"), ("curtail", "出力制御")]


def header_cols(ws):
    for row in ws.iter_rows(min_row=1, max_row=10):
        if row[0].value == "No.":
            hdr = row[0].row
            break
    else:
        raise SystemExit("'No.' ヘッダー行が見つかりません")
    cols = {}
    for c in ws[hdr]:
        h = str(c.value or "").replace("\n", "")
        cols[h] = c.column
    def find(token):
        for h, i in cols.items():
            if token in h:
                return i
        raise SystemExit(f"列 '{token}' が見つかりません")
    return hdr, {"no": find("No."), "area": find("エリア"), "name": find("変電所名"),
                 "vmax": find("最大電圧"), "vsec": find("二次電圧"), "src": find("データ出典"),
                 **{k: find(t) for k, t in FIELDS}}


def same(a, b):
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(float(a) - float(b)) < 1e-9
    # 文字列は全角半角・「♯」と「#」・空白の違いを無視する（表記ゆれで書き換えない）
    na = text(a) if a is not None else ""
    nb = text(b) if b is not None else ""
    return re.sub(r"\s", "", na) == re.sub(r"\s", "", nb)


def pick(cands, vmax, vsec, old_cap=None):
    """同名のバンクから1つに絞る。絞れなければ (None, 理由)。

    電圧（二次→一次）で絞り、それでも複数なら設備容量が本体の既存値と一致するものを選ぶ。
    設備容量は変圧器の取替えが無い限り変わらないので、前回取り込んだバンクを指す目印になる。
    """
    if len(cands) == 1:
        return cands[0], ""
    c = cands
    if isinstance(vsec, (int, float)):
        c = [x for x in c if x["v2"] is not None and abs(x["v2"] - vsec) < 0.05]
    if len(c) > 1 and isinstance(vmax, (int, float)):
        c = [x for x in c if x["v1"] is not None and abs(x["v1"] - vmax) < 0.5]
    if len(c) > 1 and isinstance(old_cap, (int, float)):
        c2 = [x for x in c if x["cap"] is not None and abs(x["cap"] - old_cap) < 0.5]
        if len(c2) == 1:
            return c2[0], ""
    if len(c) == 1:
        return c[0], ""
    # 同じ電圧のバンクが容量違いで複数並ぶことがある（同一変電所の別変圧器群）。決められないので見送る
    return None, f"同名{len(cands)}件から特定不可"


def cmd_apply(args):
    with open(os.path.join(args.cache_dir, "manifest.json"), encoding="utf-8") as f:
        manifest = json.load(f)
    pub = {}      # area -> {loose_name: [rows]}
    dates = {}    # area -> 更新日（CSV先頭行の最新）
    for area, info in manifest["areas"].items():
        for fi in info.get("files", []):
            path = os.path.join(args.cache_dir, fi["file"])
            if path.lower().endswith(".pdf"):
                from hokkaido_capacity_pdf import parse_pdf
                d, rows = parse_pdf(path)
            else:
                d, rows = parse_csv(path)
            if d:
                dates.setdefault(area, set()).add(d)
            for r in rows:
                pub.setdefault(area, {}).setdefault(loose_name(r["name"]) or r["name"], []).append(r)

    wb = load_workbook(args.src)
    ws = wb[args.list_sheet]
    hdr, C = header_cols(ws)

    # 本体側の同名（エリア＋名前）重複。重複していれば電圧で区別できるときだけ書き換える
    seen = {}
    for r in range(hdr + 1, ws.max_row + 1):
        a, n = ws.cell(r, C["area"]).value, ws.cell(r, C["name"]).value
        if a and n:
            seen.setdefault((a, loose_name(n)), []).append(r)

    report, stat = [], {}
    used = {a: set() for a in pub}
    # 照合率の分母：本体で容量データが入っている行（前回まで公表CSVと対応が取れていた行の目安）
    listed = {}
    for r in range(hdr + 1, ws.max_row + 1):
        a = ws.cell(r, C["area"]).value
        if a in pub and any(ws.cell(r, C[k]).value is not None for k in ("opcap", "cur", "up")):
            listed[a] = listed.get(a, 0) + 1
    for r in range(hdr + 1, ws.max_row + 1):
        area, name = ws.cell(r, C["area"]).value, ws.cell(r, C["name"]).value
        if area not in pub or not name:
            continue
        s = stat.setdefault(area, {"matched": 0, "changed": 0, "ambiguous": 0, "missing": 0})
        key = loose_name(name)
        cands = pub[area].get(key, [])
        vmax, vsec = ws.cell(r, C["vmax"]).value, ws.cell(r, C["vsec"]).value
        if not cands:
            s["missing"] += 1
            continue
        best, why = pick(cands, vmax, vsec, ws.cell(r, C["cap"]).value)
        if best is None and len(seen.get((area, key), [])) > 1:
            why += "（本体にも同名行あり）"
        if best is None:
            s["ambiguous"] += 1
            report.append([ws.cell(r, C["no"]).value, area, name, "特定不可", why, "", ""])
            continue
        used[area].add(id(best))
        s["matched"] += 1
        diffs = []
        for k, label in FIELDS:
            new = best[k]
            if k == "up" and not best["has_up"]:
                continue   # 上位系を公表していない会社（北陸）は既存値を触らない
            old = ws.cell(r, C[k]).value
            if new is None:
                # 公表値が「－」でも既存値は消さない。既存の予想潮流には逆算値（空容量−運用容量）や
                # 旧資料の値が入っており、空欄にすると採点材料が減るだけになるため
                continue
            if not same(old, new):
                diffs.append(f"{label}: {old if old is not None else '空'}→{new if new is not None else '空'}")
                ws.cell(r, C[k]).value = new
        if diffs:
            s["changed"] += 1
            kind = "公表PDF" if area == "北海道" else "公表CSV"
            ws.cell(r, C["src"]).value = f"{area}{kind}（{max(dates.get(area, {''}))}）"
            report.append([ws.cell(r, C["no"]).value, area, name, "更新", " / ".join(diffs), best["v1"], best["v2"]])

    # データソースシートの日付を更新
    if "データソース" in wb.sheetnames:
        ds = wb["データソース"]
        for row in ds.iter_rows(min_row=1):
            if row[0].value in dates and len(row) > 3:
                row[3].value = f"{max(dates[row[0].value])}（自動更新 {datetime.date.today().isoformat()}）"

    for area in pub:
        stat.setdefault(area, {"matched": 0, "changed": 0, "ambiguous": 0, "missing": 0})
        total = sum(len(v) for v in pub[area].values())
        stat[area]["published"] = total
        stat[area]["unlisted"] = total - len(used[area])
        stat[area]["date"] = max(dates.get(area, {""}))
        stat[area]["listed_with_data"] = listed.get(area, 0)

    wb.save(args.out)
    print("エリア | 公表日 | 公表行 | 照合 | 更新 | 特定不可 | 本体に無い公表行")
    for a, s in stat.items():
        print(f"{a} | {s['date']} | {s['published']} | {s['matched']} | {s['changed']} | {s['ambiguous']} | {s['unlisted']}")
    if args.report:
        with open(args.report, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f)
            w.writerow(["No.", "エリア", "変電所名", "結果", "変更内容・理由", "一次電圧", "二次電圧"])
            w.writerows(report)
    if args.report_md:
        write_md(args.report_md, stat, manifest)
    if args.summary_json:
        with open(args.summary_json, "w", encoding="utf-8") as f:
            json.dump({"areas": stat, "manual": MANUAL_AREAS,
                       "fetch_errors": {a: i["error"] for a, i in manifest["areas"].items() if "error" in i}},
                      f, ensure_ascii=False, indent=1)


def write_md(path, stat, manifest):
    L = ["# スコア元データ 更新レポート", "",
         f"更新: {datetime.date.today().isoformat()}（`scripts/update_capacity.py`。月1回の定期更新で自動生成）", "",
         "空容量・運用容量・予想潮流・N-1電制・出力制御を、各送配電事業者の最新の公表CSV"
         "（北海道は公表PDF「予想潮流等一覧表」）で書き換えた結果です。", "",
         "| エリア | 公表日 | 公表CSVの変電所行 | 本体と照合 | うち値が変わった | 特定不可 | 本体に無い公表行 |",
         "|---|---|---|---|---|---|---|"]
    for a, s in stat.items():
        L.append(f"| {a} | {s['date']} | {s['published']:,} | {s['matched']:,} | {s['changed']:,} | "
                 f"{s['ambiguous']:,} | {s['unlisted']:,} |")
    errs = {a: i["error"] for a, i in manifest["areas"].items() if "error" in i}
    for a, e in errs.items():
        L.append(f"| {a} | 取得失敗 | {e} | | | | |")
    L += ["", "- 「特定不可」は同じ名前の変電所が複数あり、電圧・設備容量でも1つに決められなかった行です（書き換えていません）",
          "- 「本体に無い公表行」は公表CSVにあって本体リストに無い変電所です（座標などが無いため自動では追加しません）",
          "- 公表値が「－」の項目は既存値を残しています",
          "- 北海道の「公表CSVの変電所行」はPDFから読み取ったバンク行の数です（1変電所に複数バンクがあれば複数行）", "",
          "## 自動では更新できないエリア（手作業で確認）", "",
          "| エリア | 理由・確認方法 |", "|---|---|"]
    for a, why in MANUAL_AREAS.items():
        L.append(f"| {a} | {why} |")
    L.append("")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(L))


def main():
    ap = argparse.ArgumentParser(description="公表CSVで変電所の空容量などを更新する")
    sub = ap.add_subparsers(dest="cmd", required=True)
    f = sub.add_parser("fetch", help="各社の最新CSVをダウンロード（要ネットワーク）")
    f.add_argument("--cache-dir", required=True)
    f.add_argument("--sleep", type=float, default=1.0)
    f.add_argument("--areas", nargs="*", help="取得するエリアを限定する（例: --areas 北海道）。省略時は全エリア")
    f.set_defaults(func=cmd_fetch)
    a = sub.add_parser("apply", help="本体xlsxの容量欄を書き換える")
    a.add_argument("--in", dest="src", required=True)
    a.add_argument("--cache-dir", required=True)
    a.add_argument("--out", required=True)
    a.add_argument("--list-sheet", default="全国変電所リスト")
    a.add_argument("--report")
    a.add_argument("--summary-json")
    a.add_argument("--report-md", help="更新レポート（Markdown）の出力先")
    a.set_defaults(func=cmd_apply)
    args = ap.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
