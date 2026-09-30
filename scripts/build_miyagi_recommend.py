"""宮城県の推奨変電所リストを作る（社外提示用＝NDA前 / 社内用 の2ファイル）

  python scripts/build_miyagi_recommend.py --addr-cache miyagi_addr.json

- 元データ: data/substations.js（変電所スコアリング v1）
- 所在地: 国土地理院 逆ジオコーダで取得した市区町村コード・町名（--addr-cache に保存）
- 立地区分は住所からの机上推定（現地未確認）。下の LOCATION を手で更新する。

社外提示用には公表情報（変電所名・所在地・電圧・公表空容量）と当社のおすすめ度だけを載せ、
社内スコアの配点・接続検討の実績は載せない（NDA締結前のため）。
緯度経度は載せる（変電所の位置は公開情報。東北電力NWも県別の系統状況マップPDFで概略位置を公表）。
"""
import argparse
import json
import math
import time
import urllib.request

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

MUNI = {
    "04101": "仙台市青葉区", "04102": "仙台市宮城野区", "04103": "仙台市若林区",
    "04104": "仙台市太白区", "04105": "仙台市泉区", "04202": "石巻市", "04203": "塩竈市",
    "04207": "名取市", "04208": "角田市", "04209": "多賀城市", "04211": "岩沼市",
    "04212": "登米市", "04213": "栗原市", "04214": "東松島市", "04215": "大崎市",
    "04216": "富谷市", "04322": "村田町", "04323": "柴田町", "04341": "丸森町",
    "04361": "亘理町", "04401": "松島町", "04406": "利府町", "04421": "大和町",
    "04424": "大衡村", "04445": "加美町", "04501": "涌谷町", "04505": "美里町",
    "04581": "女川町", "04606": "南三陸町",
}

# 立地区分（住所からの机上推定）と社外向けコメント。tier: ◎ が社外提示対象
LOCATION = {
    "高舘変電所": ("◎", "郊外（農地・丘陵）", "空容量に余裕。名取市西部で周辺に農地が多く、用地を探しやすい見込み"),
    "中山変電所": ("◎", "郊外（青葉区西部）", "154kV受電・空容量に余裕。仙台市西部の郊外。市街化調整区域の開発許可の扱いは要確認"),
    "沼辺変電所": ("◎", "農村部", "空容量に余裕。周辺は農地中心で用地を探しやすい見込み"),
    "陸前大幡変電所": ("◎", "郊外（古川地区）", "空容量に余裕。大崎市古川の郊外"),
    "利府変電所": ("◎", "郊外", "空容量に余裕。仙台近郊でアクセス良好"),
    "美田園変電所": ("◎", "郊外（仙台空港近く）", "空容量に余裕。沿岸部のため津波浸水想定の確認が必要"),
    "富谷変電所": ("◎", "郊外", "空容量に余裕。仙台北部の郊外"),
    "玉浦変電所": ("◎", "沿岸の農地", "空容量あり。周辺は農地中心。津波浸水想定の確認が必要"),
    "東北電力株式会社 岩沼変電所": ("○", "市街地寄り", "空容量に余裕。周辺は市街地のため用地は郊外側で探す想定"),
    "陸前山下変電所": ("○", "市街地寄り", "空容量あり。石巻市西部"),
    "船岡変電所": ("○", "市街地寄り", "空容量あり。柴田町"),
    "成田変電所": ("○", "住宅地", "空容量あり。周辺は住宅地のため用地確保に工夫が必要"),
    "吉岡変電所": ("○", "郊外", "空容量あり。大和町中心部の近く"),
    "増田変電所": ("○", "市街地寄り", "空容量あり。名取市中心部の近く"),
    "蛇田変電所": ("○", "住宅地", "空容量あり。石巻市の住宅地"),
    "佐沼変電所": ("○", "市街地寄り", "空容量あり。登米市中心部の近く"),
    "涌谷変電所": ("○", "農村部", "空容量あり。周辺は農地中心"),
    "福田町変電所": ("○", "工業地", "空容量が大きい。工業地で用途地域面は有利だが用地単価は高め"),
    "重吉変電所": ("○", "港湾・工業地", "空容量が大きい。石巻港周辺。津波浸水想定の確認が必要"),
    "卸町変電所": ("○", "流通業務地", "空容量が大きい。用地単価は高め"),
    "新浜変電所": ("○", "港湾・工業地", "空容量が大きい。塩釜港周辺。津波浸水想定の確認が必要"),
    "明通変電所": ("○", "市街地（泉区）", "154kV受電・空容量が大きい。用地単価は高め"),
    "石巻湊変電所": ("○", "港湾・市街地", "空容量が大きい。津波浸水想定の確認が必要"),
    "塩釜築港変電所": ("○", "港湾・工業地", "空容量あり。津波浸水想定の確認が必要"),
    # 以下は社内のみ（△）
    "宮崎変電所": ("△", "農村部", "空容量が小さめ(9MW)"),
    "中新田変電所": ("△", "郊外", "空容量が小さめ(9MW)"),
    "南二又変電所": ("△", "農村部", "空容量が小さめ(9MW)"),
    "飯野川変電所": ("△", "農村部", "空容量が小さめ(9MW)"),
    "鬼首変電所": ("△", "山間部（温泉地）", "空容量が小さい(6MW)・山間部で輸送も要確認"),
    "鮎川変電所": ("△", "半島部", "空容量が小さい(7MW)・半島先端で輸送も要確認"),
    "女川変電所": ("△", "沿岸部", "空容量あり(10MW)だが平地が少ない"),
    "松島変電所": ("△", "観光地", "空容量が小さめ(7MW)・景観配慮"),
    "登米変電所": ("△", "農村部", "空容量が小さめ(7MW)"),
    "豊里変電所": ("△", "農村部", "空容量が小さい(6MW)"),
    "小野変電所": ("△", "農村部", "空容量が小さい(6MW)"),
    "松の平変電所": ("△", "住宅地", "空容量が小さめ(7MW)"),
    "松坂変電所": ("△", "郊外", "空容量が小さめ(7MW)"),
    "根白石変電所": ("△", "郊外", "空容量が小さめ(9MW)"),
    "高森変電所": ("△", "住宅地", "空容量が小さめ(8MW)"),
    "亘理変電所": ("△", "市街地寄り", "空容量が小さめ(9MW)"),
    "小牛田変電所": ("△", "市街地寄り", "空容量が小さめ(8MW)"),
    "志津川変電所": ("△", "沿岸部", "空容量が小さめ(7MW)"),
    "角田変電所": ("△", "市街地寄り", "空容量が小さめ(7MW)"),
    "丸森変電所": ("△", "農村部", "空容量が小さい(6MW)"),
    "渡波変電所": ("△", "沿岸部", "空容量あり(10MW)・津波浸水想定の確認が必要"),
}
# 仙台都心・市街地の中心部（空容量はあるが蓄電所用地が見つからない）
URBAN = "都心・市街地中心（用地確保が困難）"
URBAN_NAMES = {
    "仙台本町変電所", "榴岡変電所", "勾当台変電所", "広瀬通変電所", "柳町通変電所",
    "堤通変電所", "大町変電所", "東北電力花京院変電所", "土樋変電所", "鍋田変電所",
    "あすと長町変電所", "長町変電所", "八木山変電所", "東北電力ネットワーク 苦竹変電所",
    "東北電力ネットワーク 南小泉変電所", "荒巻変電所", "西多賀変電所", "中田変電所",
    "鶴ヶ谷変電所", "小鶴新田変電所", "七北田変電所", "沖野変電所", "大日変電所",
    "多賀城変電所", "塩釜変電所", "泉変電所", "旭ヶ丘変電所", "大富変電所",
}


def load_rows(path):
    s = open(path, encoding="utf-8").read()
    d = json.loads(s[s.index("{"): s.rstrip().rstrip(";").rindex("}") + 1])
    rows = []
    for r in d["rows"]:
        r = dict(zip(d["fields"], r))
        if r["pref"] != "宮城県":
            continue
        for k in ("area", "acc", "n1", "curtail", "rank"):
            r[k] = d["dict"][k][r[k]]
        rows.append(r)
    return rows, d["meta"]["generated"]


def osm_check(rows, osm_path):
    """座標から100m以内にOSMの変電所があれば位置確認済みとする（名称一致でない座標の検証用）"""
    pts = []
    for e in json.load(open(osm_path, encoding="utf-8"))["elements"]:
        c = e if "lat" in e else e.get("center")
        if c:
            pts.append((c["lat"], c["lon"]))
    for r in rows:
        r["pos_ok"] = None
        if r["lat"] is None:
            continue
        if r["acc"].startswith("OSM"):
            r["pos_ok"] = True
            continue
        d = min(6371000 * math.hypot(math.radians(p[0] - r["lat"]),
                                     math.radians(p[1] - r["lon"]) * math.cos(math.radians(r["lat"])))
                for p in pts)
        r["pos_ok"] = d <= 100


def reverse_geocode(rows, cache_path):
    try:
        cache = json.load(open(cache_path, encoding="utf-8"))
    except FileNotFoundError:
        cache = {}
    for r in rows:
        if r["lat"] is None:
            continue
        key = f"{r['lat']},{r['lon']}"
        if key not in cache:
            url = ("https://mreversegeocoder.gsi.go.jp/reverse-geocoder/LonLatToAddress"
                   f"?lat={r['lat']}&lon={r['lon']}")
            res = json.load(urllib.request.urlopen(url, timeout=15))["results"]
            cache[key] = [res["muniCd"], res["lv01Nm"]]
            time.sleep(0.2)
        code, town = cache[key]
        r["addr_muni"] = MUNI.get(code, code)
        r["addr_town"] = town
    json.dump(cache, open(cache_path, "w", encoding="utf-8"), ensure_ascii=False, indent=0)


HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
TIER_FILL = {"◎": "E2EFDA", "○": "FFF2CC", "△": "F2F2F2"}
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def write_table(ws, start_row, headers, rows, widths, tier_col=None):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(start_row, c, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    for i, row in enumerate(rows, start_row + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(i, c, v)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if tier_col is not None and row[tier_col] in TIER_FILL:
                cell.fill = PatternFill("solid", fgColor=TIER_FILL[row[tier_col]])
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(start_row, c).column_letter].width = w
    ws.freeze_panes = ws.cell(start_row + 1, 1)


def addr(r):
    if r.get("addr_muni"):
        return f"{r['addr_muni']}{r['addr_town'] or ''}"
    return "（未取得）"


def build_external(rows, out, asof, today):
    wb = Workbook()
    ws = wb.active
    ws.title = "宮城県 推奨変電所"
    ws["A1"] = "宮城県 系統用蓄電池（高圧・2MW級）向け 推奨変電所リスト"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"作成日 {today} ／ 空容量は東北電力ネットワークの公表値（当社調べ・{asof}時点）"
    ws["A3"] = "掲載基準：空容量に余裕があり、郊外で蓄電所用地を探しやすい変電所"
    for a in ("A2", "A3"):
        ws[a].font = Font(size=9, color="595959")
    picked = [r for r in rows if LOCATION.get(r["name"], ("",))[0] == "◎"]
    picked.sort(key=lambda r: -(r["availCur"] or 0))
    table = []
    for n, r in enumerate(picked, 1):
        tier, loc, note = LOCATION[r["name"]]
        table.append([n, tier, r["name"].replace("東北電力株式会社 ", ""), addr(r),
                      f"{r['vmax']}kV" if r["vmax"] else "－", r["availCur"], loc, note,
                      round(r["lat"], 5), round(r["lon"], 5), "地図で開く",
                      "確認済" if r["pos_ok"] else "目安（要現地確認）"])
    write_table(ws, 5, ["No", "おすすめ度", "変電所名", "所在地", "最大電圧", "公表空容量\n(MW)", "立地", "ひとこと",
                        "緯度", "経度", "地図", "位置"],
                table, [5, 9, 20, 30, 9, 11, 20, 60, 10, 11, 11, 18], tier_col=1)
    for i, r in enumerate(picked, 6):
        cell = ws.cell(i, 11)
        cell.hyperlink = f"https://www.google.com/maps?q={r['lat']},{r['lon']}"
        cell.font = Font(color="0563C1", underline="single")
    end = 5 + len(table) + 2
    notes = [
        "【ご留意事項】",
        "・本リストは公表情報をもとにした机上の一次スクリーニングです。連系の可否・負担金・工期は一般送配電事業者の接続検討で確定します。",
        "・空容量は公表時点の値で、先行申込により変動します。最新値は東北電力ネットワークの空容量マップでご確認ください。",
        "・宮城県内の該当変電所は、いずれもN-1電制「不可」・出力制御「あり」の公表条件です。",
        "・立地区分は所在地からの推定で、現地・用途地域・ハザード（津波浸水等）は未確認です。",
        "・緯度経度は地図データ（OpenStreetMap等）から当社が取得した変電所の位置です。「確認済」は地図上の変電所設備と一致したものです。"
        + ("「目安」は名称検索による位置で数百m〜数kmずれる可能性があります。" if not all(r["pos_ok"] for r in picked) else ""),
        "・具体的な候補地・接続検討の実績等の詳細は、秘密保持契約の締結後にご案内いたします。",
    ]
    for i, t in enumerate(notes):
        ws.cell(end + i, 1, t).font = Font(size=9, bold=(i == 0))
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0
    wb.save(out)
    return len(table)


def build_internal(rows, out, asof, today):
    wb = Workbook()
    ws = wb.active
    ws.title = "宮城県 全変電所"
    ws["A1"] = f"宮城県 変電所一覧（社内用・社外秘） 作成 {today} ／ スコアリングv1・{asof}データ"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "社外提示=◎のみ（別ファイル）。○=条件付き候補（社内のみ）。△=空容量小さめ/条件悪。都心=空容量はあるが用地確保困難。スコア・実績は社外に出さないこと（座標は社外用にも掲載）。"
    ws["A2"].font = Font(size=9, color="C00000")

    def cls(r):
        if r["name"] in LOCATION:
            return LOCATION[r["name"]]
        if r["name"] in URBAN_NAMES:
            return ("都心", URBAN, "")
        if r["rank"] == "除外":
            return ("対象外", "", "ゲート除外（空容量0・上位系増強リスク）")
        if r["rank"] == "データ無":
            return ("対象外", "", "空容量が未公表（基幹系・開閉所など）")
        if r["rank"] == "C":
            return ("対象外", "", "空容量が小さい")
        return ("対象外", "", "座標未取得・所在地不明（番号名の変電所など）")

    order = {"◎": 0, "○": 1, "△": 2, "都心": 3, "対象外": 4}
    rows = sorted(rows, key=lambda r: (order[cls(r)[0]], -(r["availCur"] or -1), -(r["grid"] or 0)))
    table = []
    for r in rows:
        tier, loc, note = cls(r)
        table.append([tier, r["name"], addr(r), r["vmax"], r["vsec"], r["availCur"], r["opcap"], r["flow"],
                      r["n1"], r["curtail"], r["grid"], r["rank"], loc, note,
                      r["lat"], r["lon"], r["acc"]])
    write_table(ws, 4, ["区分", "変電所名", "所在地", "最大電圧\n(kV)", "二次電圧\n(kV)", "空容量\n(MW)",
                        "運用容量\n(MW)", "予想潮流\n(MW)", "N-1電制", "出力制御", "系統スコア\n(/80)",
                        "系統ランク", "立地（机上推定）", "メモ", "緯度", "経度", "座標精度"],
                table, [7, 26, 30, 8, 8, 8, 8, 8, 10, 8, 9, 8, 24, 44, 10, 11, 20], tier_col=0)
    ws.auto_filter.ref = f"A4:Q{4 + len(table)}"
    wb.save(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--data", default="data/substations.js")
    ap.add_argument("--osm", default="osm_cache/宮城県.json")
    ap.add_argument("--addr-cache", default="geo_cache/miyagi_addr.json")
    ap.add_argument("--today", default=time.strftime("%Y-%m-%d"))
    ap.add_argument("--out-ext", default="docs/宮城県_推奨変電所リスト_社外提示用(NDA前).xlsx")
    ap.add_argument("--out-int", default="docs/宮城県_変電所リスト_社内用.xlsx")
    a = ap.parse_args()
    rows, asof = load_rows(a.data)
    reverse_geocode(rows, a.addr_cache)
    osm_check(rows, a.osm)
    n = build_external(rows, a.out_ext, asof, a.today)
    build_internal(rows, a.out_int, asof, a.today)
    print(f"社外提示用 {n}件 → {a.out_ext}\n社内用 {len(rows)}件 → {a.out_int}")


if __name__ == "__main__":
    main()
